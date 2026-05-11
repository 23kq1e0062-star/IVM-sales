from flask import Flask, request, jsonify, send_from_directory, send_file
from flask_cors import CORS
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
import sqlite3, hashlib, io, re, os
from datetime import date, datetime

app = Flask(__name__, static_folder='.')
CORS(app)

SUPPLIERS   = ['A K ENTERPRISES', 'BHAGAVATHI FANCY', 'SRM SONS']
DEPARTMENTS = ['ACCESSORIES', 'FOOTWEAR', 'HOME NEEDS', 'LIFESTYLE & FASHION',
               'NEWBORN', 'PERSONAL CARE & BEAUTY', 'TEXTILES']

DB_PATH     = os.path.join(os.path.dirname(__file__), 'sale_data.db')
MASTER_PASS = hashlib.sha256(b'admin123').hexdigest()

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    with get_db() as conn:
        conn.executescript('''
            CREATE TABLE IF NOT EXISTS master_items (
                item_code TEXT PRIMARY KEY,
                item_name TEXT, department TEXT,
                supplier TEXT, updated_at TEXT
            );
            CREATE TABLE IF NOT EXISTS daily_sales (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                sale_date TEXT, month_key TEXT,
                item_code TEXT, item_name TEXT,
                department TEXT, supplier TEXT,
                sold_value REAL, cost_value REAL, created_at TEXT
            );
            CREATE TABLE IF NOT EXISTS app_config (
                key TEXT PRIMARY KEY, value TEXT
            );
        ''')
        conn.execute("INSERT OR IGNORE INTO app_config VALUES ('master_password',?)", (MASTER_PASS,))

init_db()

def norm(v): return str(v).strip().upper() if pd.notna(v) else ''

def clean_code(v):
    if not pd.notna(v): return ''
    try: return str(int(float(str(v).strip())))
    except: return str(v).strip()

def check_password(pwd):
    with get_db() as conn:
        row = conn.execute("SELECT value FROM app_config WHERE key='master_password'").fetchone()
        return row and row['value'] == hashlib.sha256(pwd.encode()).hexdigest()

@app.route('/')
def index(): return send_from_directory('.', 'index.html')

@app.route('/change-password', methods=['POST'])
def change_password():
    d = request.get_json()
    if not check_password(d.get('current','')):
        return jsonify({'error': 'Incorrect current password'}), 403
    new_hash = hashlib.sha256(d.get('new','').encode()).hexdigest()
    with get_db() as conn:
        conn.execute("UPDATE app_config SET value=? WHERE key='master_password'", (new_hash,))
    return jsonify({'success': True})

@app.route('/upload-master', methods=['POST'])
def upload_master():
    pwd = request.form.get('password','')
    if not check_password(pwd):
        return jsonify({'error': 'Incorrect password'}), 403
    if 'file' not in request.files:
        return jsonify({'error': 'No file'}), 400
    try:
        df = pd.read_excel(request.files['file'], sheet_name=0)
        df.columns = [str(c).strip().upper() for c in df.columns]
        cc = next((c for c in df.columns if 'ITEM CODE' in c or c=='CODE'), df.columns[0])
        dc = next((c for c in df.columns if 'DEPARTMENT' in c), None)
        nc = next((c for c in df.columns if 'ITEM NAME' in c), None)
        sc = next((c for c in df.columns if 'SUPPLIER' in c), None)
        now = datetime.now().isoformat()
        with get_db() as conn:
            conn.execute("DELETE FROM master_items")
            for _, row in df.iterrows():
                code = clean_code(row[cc])
                if not code or code=='nan': continue
                conn.execute("INSERT OR REPLACE INTO master_items VALUES (?,?,?,?,?)",
                    (code,
                     str(row[nc]).strip() if nc and pd.notna(row[nc]) else '',
                     str(row[dc]).strip().upper() if dc and pd.notna(row[dc]) else '',
                     str(row[sc]).strip().upper() if sc and pd.notna(row[sc]) else '',
                     now))
        count = conn.execute("SELECT COUNT(*) FROM master_items").fetchone()[0]
        return jsonify({'success': True, 'count': count})
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.route('/master-status')
def master_status():
    with get_db() as conn:
        row = conn.execute("SELECT COUNT(*) as cnt, MAX(updated_at) as updated FROM master_items").fetchone()
        return jsonify({'count': row['cnt'], 'updated': row['updated']})

@app.route('/upload', methods=['POST'])
def upload():
    if 'file' not in request.files:
        return jsonify({'error': 'No file uploaded'}), 400
    file    = request.files['file']
    walkins = int(request.form.get('walkins','0') or 0)
    manual_date = request.form.get('sale_date','').strip()  # DD-MM-YYYY from user

    try:
        xls         = pd.ExcelFile(file)
        sheet_names = xls.sheet_names
    except Exception as e:
        return jsonify({'error': f'Failed to read Excel: {e}'}), 400

    # Find item-level sheet
    copy_sheet = next((s for s in sheet_names if s.upper()=='COPY'), None)
    copy_df = None
    if copy_sheet:
        copy_df = pd.read_excel(xls, sheet_name=copy_sheet)
        copy_df.columns = [str(c).strip().upper() for c in copy_df.columns]
    else:
        for sname in sheet_names:
            try:
                raw = pd.read_excel(xls, sheet_name=sname, header=None)
                hr  = next((i for i, row in raw.iterrows()
                            if any(norm(str(v)) in ('ITEM CODE','ITEM NAME') for v in row if pd.notna(v))), None)
                if hr is not None:
                    df_try = pd.read_excel(xls, sheet_name=sname, header=hr)
                    df_try.columns = [str(c).strip().upper() for c in df_try.columns]
                    if any('TOTAL SOLD VALUE' in c for c in df_try.columns):
                        copy_df = df_try; break
            except: continue

    if copy_df is None:
        return jsonify({'error': f'Could not find item detail data. Sheets: {sheet_names}'}), 400

    ic_col   = next((c for c in copy_df.columns if 'ITEM CODE' in c), None)
    dept_col = next((c for c in copy_df.columns if 'DEPARTMENT' in c), None)
    sup_col  = next((c for c in copy_df.columns if 'SUPPLIER NAME' in c), None)
    tsv_col  = next((c for c in copy_df.columns if 'TOTAL SOLD VALUE' in c), None)
    tcv_col  = next((c for c in copy_df.columns if 'TOTAL COST VALUE' in c), None)
    name_col = next((c for c in copy_df.columns if 'ITEM NAME' in c), None)
    bdt_col  = next((c for c in copy_df.columns if 'BILL DT' in c or 'BILL DATE' in c), None)

    if not tsv_col:
        return jsonify({'error': f'Total Sold Value column not found. Columns: {list(copy_df.columns)}'}), 400

    if ic_col:
        copy_df = copy_df[copy_df[ic_col].notna()]
        copy_df = copy_df[copy_df[ic_col].astype(str).str.strip().str.lower() != 'nan']
    copy_df[tsv_col] = pd.to_numeric(copy_df[tsv_col], errors='coerce').fillna(0)

    # Load master
    with get_db() as conn:
        master_rows = conn.execute("SELECT item_code, department FROM master_items").fetchall()
        def nc2(c):
            try: return str(int(float(str(c).strip())))
            except: return str(c).strip()
        master_map = {nc2(r['item_code']): r['department'] for r in master_rows}

    # Bill Dt map
    bill_dt_map = {}
    if bdt_col:
        for _, row in copy_df.iterrows():
            code = clean_code(row[ic_col]) if ic_col else ''
            dt   = row[bdt_col]
            if code and code != 'nan':
                if hasattr(dt, 'strftime'): dt = dt.strftime('%d-%m-%Y')
                bill_dt_map[code] = str(dt) if pd.notna(dt) else ''

    # Determine sale_date: manual input > bill dt > today
    if manual_date and re.match(r'\d{2}-\d{2}-\d{4}', manual_date):
        sale_date = manual_date
    elif bill_dt_map:
        first_dt = next(iter(bill_dt_map.values()), '')
        sale_date = first_dt if re.match(r'\d{2}-\d{2}-\d{4}', first_dt) else date.today().strftime('%d-%m-%Y')
    else:
        sale_date = date.today().strftime('%d-%m-%Y')
    month_key = sale_date[3:]  # MM-YYYY

    # Calculate totals
    sup_totals  = {s: {'value':0.0,'qty':0} for s in SUPPLIERS}
    dept_totals = {d: {'value':0.0,'qty':0} for d in DEPARTMENTS}
    total_value = 0.0; total_qty = 0
    detail_rows = []; db_rows = []

    for _, row in copy_df.iterrows():
        code  = clean_code(row[ic_col]) if ic_col else ''
        if not code or code == 'nan': continue
        sup   = str(row[sup_col]).strip().upper() if sup_col and pd.notna(row[sup_col]) else ''
        dept  = master_map.get(code,'') or (str(row[dept_col]).strip().upper() if dept_col and pd.notna(row[dept_col]) else '')
        val   = float(row[tsv_col])
        cost  = float(row[tcv_col]) if tcv_col and pd.notna(row[tcv_col]) else 0.0
        iname = str(row[name_col]).strip() if name_col and pd.notna(row[name_col]) else ''
        bdt   = bill_dt_map.get(code,'')

        total_value += val; total_qty += 1
        if sup in sup_totals:
            sup_totals[sup]['value'] += val; sup_totals[sup]['qty'] += 1
        if dept in dept_totals:
            dept_totals[dept]['value'] += val; dept_totals[dept]['qty'] += 1

        detail_rows.append({'item_code':code,'department':dept,'item_name':iname,
                            'supplier':sup,'sold_value':val,'cost_value':cost,'bill_dt':bdt})
        db_rows.append((sale_date,month_key,code,iname,dept,sup,val,cost,datetime.now().isoformat()))

    with get_db() as conn:
        conn.execute("DELETE FROM daily_sales WHERE sale_date=?", (sale_date,))
        conn.executemany(
            "INSERT INTO daily_sales (sale_date,month_key,item_code,item_name,department,supplier,sold_value,cost_value,created_at) VALUES (?,?,?,?,?,?,?,?,?)",
            db_rows)

    return jsonify({
        'suppliers':   [{'name':s.title(),'value':round(v['value'],2),'qty':v['qty']} for s,v in sup_totals.items()],
        'departments': [{'name':d.title(),'value':round(v['value'],2),'qty':v['qty']} for d,v in dept_totals.items()],
        'total_value': round(total_value,2), 'total_qty': total_qty,
        'walkins': walkins, 'sale_date': sale_date,
        'detail_rows': detail_rows, 'has_master': len(master_map)>0
    })

@app.route('/monthly-report', methods=['GET'])
def monthly_report():
    month = request.args.get('month','')
    with get_db() as conn:
        if not month:
            months = conn.execute("SELECT DISTINCT month_key FROM daily_sales ORDER BY month_key").fetchall()
            return jsonify({'months': [r['month_key'] for r in months]})
        rows = conn.execute("SELECT * FROM daily_sales WHERE month_key=?", (month,)).fetchall()
        if not rows:
            return jsonify({'error': 'No data for this month'}), 404
        dept_totals = {}; sup_totals = {}; item_totals = {}
        for r in rows:
            d = r['department'] or 'OTHER'
            dept_totals[d] = dept_totals.get(d,{'value':0.0,'qty':0})
            dept_totals[d]['value'] += r['sold_value']; dept_totals[d]['qty'] += 1
            s = r['supplier'] or 'OTHER'
            sup_totals[s] = sup_totals.get(s,{'value':0.0,'qty':0})
            sup_totals[s]['value'] += r['sold_value']; sup_totals[s]['qty'] += 1
            key = r['item_code']
            item_totals[key] = item_totals.get(key,{'name':r['item_name'],'dept':r['department'],'value':0.0,'qty':0})
            item_totals[key]['value'] += r['sold_value']; item_totals[key]['qty'] += 1
        items_sorted = sorted(item_totals.items(), key=lambda x: x[1]['value'], reverse=True)
        top10    = [{'code':k,'name':v['name'],'dept':v['dept'],'value':round(v['value'],2),'qty':v['qty']} for k,v in items_sorted[:10]]
        bottom10 = [{'code':k,'name':v['name'],'dept':v['dept'],'value':round(v['value'],2),'qty':v['qty']} for k,v in items_sorted[-10:]]
        return jsonify({
            'month': month,
            'departments': [{'name':k,'value':round(v['value'],2),'qty':v['qty']} for k,v in sorted(dept_totals.items(),key=lambda x:-x[1]['value'])],
            'suppliers':   [{'name':k,'value':round(v['value'],2),'qty':v['qty']} for k,v in sorted(sup_totals.items(), key=lambda x:-x[1]['value'])],
            'top_items': top10, 'bottom_items': bottom10,
            'total_value': round(sum(r['sold_value'] for r in rows),2),
            'total_qty': len(rows)
        })

@app.route('/clear-data', methods=['POST'])
def clear_data():
    d = request.get_json()
    scope = d.get('scope', 'all')  # 'all', 'date', 'month'
    target = d.get('target', '')
    with get_db() as conn:
        if scope == 'all':
            conn.execute("DELETE FROM daily_sales")
        elif scope == 'date' and target:
            conn.execute("DELETE FROM daily_sales WHERE sale_date=?", (target,))
        elif scope == 'month' and target:
            conn.execute("DELETE FROM daily_sales WHERE month_key=?", (target,))
        count = conn.execute("SELECT COUNT(*) FROM daily_sales").fetchone()[0]
    return jsonify({'success': True, 'remaining': count})

@app.route('/download-excel', methods=['POST'])
def download_excel():
    data        = request.get_json()
    detail_rows = data.get('detail_rows',[])
    summary     = data.get('summary',{})
    walkins     = data.get('walkins',0)
    sale_date   = data.get('sale_date', date.today().strftime('%d-%m-%Y'))
    title       = f'{sale_date} SALE DATA'

    wb   = openpyxl.Workbook()
    thin = Side(style='thin', color='C5D4F0')
    bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)

    def st(cell, bold=False, fg='000000', bg=None, align='left', fmt=None, sz=10):
        cell.font      = Font(name='Calibri', bold=bold, color=fg, size=sz)
        if bg: cell.fill = PatternFill('solid', start_color=bg)
        cell.alignment = Alignment(horizontal=align, vertical='center')
        cell.border    = bdr
        if fmt: cell.number_format = fmt

    def title_row(ws, text, cols, bg='1A3C8F'):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
        c = ws.cell(row=1, column=1, value=text)
        st(c, bold=True, fg='FFFFFF', bg=bg, align='center', sz=13)
        ws.row_dimensions[1].height = 24

    def hdr_row(ws, headers, row=2, bg='2756C5'):
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=row, column=ci, value=h)
            st(c, bold=True, fg='FFFFFF', bg=bg, align='center', sz=10)
        ws.row_dimensions[row].height = 18

    def clean_code(v):
        try: return int(float(str(v)))
        except: return str(v)

    # Aggregate items — composite key (item_code + supplier) so each supplier tracked separately
    item_map = {}
    for item in detail_rows:
        code = str(clean_code(item.get('item_code','')))
        sup  = item.get('supplier','').upper().strip()
        k    = f"{code}||{sup}"
        if k not in item_map:
            item_map[k] = {'code': clean_code(item.get('item_code','')),
                           'name': item.get('item_name','').upper(),
                           'dept': item.get('department','').title(),
                           'supplier': sup,
                           'qty': 0, 'value': 0.0, 'cost': 0.0}
        item_map[k]['qty']   += 1
        item_map[k]['value'] += item.get('sold_value', 0)
        item_map[k]['cost']  += item.get('cost_value', 0)

    all_items = list(item_map.values())
    by_qty    = sorted(all_items, key=lambda x: x['qty'],   reverse=True)
    by_value  = sorted(all_items, key=lambda x: x['value'], reverse=True)

    # ── Sheet 1: Item Detail ──────────────────────────────
    ws1 = wb.active; ws1.title = 'Item Detail'
    ws1.freeze_panes = 'A3'
    title_row(ws1, title, 9)
    hdr_row(ws1, ['Item Code','Department','Item Name','Supplier Name','Sold Value (Rs)','Cost Value (Rs)','Profit (Rs)','Profit %','Bill Date'])
    for ri, item in enumerate(detail_rows, 3):
        bg = 'F0F4FF' if ri%2==0 else None
        sold  = item.get('sold_value', 0)
        cost  = item.get('cost_value', 0)
        profit = round(sold - cost, 2)
        profit_pct = round((profit / sold * 100), 2) if sold else 0.0
        # Colour profit cell green if positive, red if negative
        profit_fg = '1A7A3A' if profit >= 0 else 'C0392B'
        vals = [clean_code(item.get('item_code','')), item.get('department','').title(),
                item.get('item_name','').upper(), item.get('supplier','').upper(),
                sold, cost, profit, profit_pct, item.get('bill_dt','')]
        for ci, v in enumerate(vals, 1):
            c = ws1.cell(row=ri, column=ci, value=v)
            # Use profit colour for profit and profit% columns
            fg = profit_fg if ci in (7, 8) else '1A2340'
            bold = ci in (7, 8)
            st(c, bold=bold, fg=fg, bg=bg,
               align='right' if ci in (5,6,7,8) else 'center' if ci in (1,9) else 'left',
               fmt='#,##0.00' if ci in (5,6,7) else '0.00"%"' if ci==8 else '0' if ci==1 else None, sz=9)
    for col,w in zip('ABCDEFGHI',[11,24,36,20,15,15,14,11,13]): ws1.column_dimensions[col].width=w

    # ── Sheet 2: Sale Report ──────────────────────────────
    ws2 = wb.create_sheet('Sale Report')
    title_row(ws2, title, 3)
    hdr_row(ws2, ['Vendor / Department','Total Sold Value (Rs)','Sold Items'])
    r = 3
    def sec(label):
        nonlocal r
        ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        st(ws2.cell(row=r, column=1, value=label), bold=True, fg='1A3C8F', bg='D0DDF7', sz=9); r+=1
    def drow(name, value, qty):
        nonlocal r
        bg = 'F0F4FF' if r%2==0 else None
        for ci,v in enumerate([name,value,qty],1):
            c=ws2.cell(row=r,column=ci,value=v)
            st(c,fg='1A2340',bg=bg,align='right' if ci>1 else 'left',
               fmt='#,##0.00' if ci==2 else '#,##0' if ci==3 else None,sz=10)
        r+=1
    sec('Suppliers')
    for s in summary.get('suppliers',[]): drow(s['name'],s['value'],s['qty'])
    for d in summary.get('departments',[]): drow(d['name'],d['value'],d['qty'])
    for ci,v in enumerate(['TOTAL',summary.get('total_value',0),summary.get('total_qty',0)],1):
        c=ws2.cell(row=r,column=ci,value=v)
        st(c,bold=True,fg='FFFFFF',bg='1A3C8F',align='right' if ci>1 else 'left',
           fmt='#,##0.00' if ci==2 else '#,##0' if ci==3 else None,sz=11)
    r+=1
    ws2.merge_cells(f'B{r}:C{r}')
    for ci,v in enumerate(['WALK IN',walkins],1):
        c=ws2.cell(row=r,column=ci,value=v)
        st(c,bold=True,fg='1A7A3A',bg='E8F7EC',align='center' if ci==2 else 'left',fmt='#,##0' if ci==2 else None,sz=10)
    ws2.column_dimensions['A'].width=34; ws2.column_dimensions['B'].width=22; ws2.column_dimensions['C'].width=14

    # ── Sheet 3: Qty Ranking ──────────────────────────────
    ws3 = wb.create_sheet('Qty Ranking')
    ws3.freeze_panes = 'A3'
    title_row(ws3, f'{title} — Items Ranked by Qty Sold (High to Low)', 7, bg='1A7A3A')
    hdr_row(ws3, ['Rank','Item Code','Item Name','Department','Supplier','Qty Sold','Total Value (Rs)'], bg='1A7A3A')
    for ri, item in enumerate(by_qty, 3):
        bg = 'F0FFF4' if ri%2==0 else 'FFFFFF'
        vals = [ri-2, item['code'], item['name'], item['dept'], item['supplier'], item['qty'], item['value']]
        for ci,v in enumerate(vals,1):
            c=ws3.cell(row=ri,column=ci,value=v)
            st(c,bold=(ci==1),fg='155A2C' if ci in (1,6) else '1A2340',bg=bg,
               align='center' if ci in (1,2) else 'right' if ci in (6,7) else 'left',
               fmt='#,##0.00' if ci==7 else '#,##0' if ci==6 else '0' if ci==2 else None,sz=9)
    for col,w in zip('ABCDEFG',[7,11,36,24,20,11,16]): ws3.column_dimensions[col].width=w
    # Chart
    top15q = by_qty[:min(15,len(by_qty))]
    if top15q:
        ws3.cell(row=2,column=9,value='Item'); ws3.cell(row=2,column=10,value='Qty')
        for i,item in enumerate(top15q,3):
            ws3.cell(row=i,column=9,value=(item['name'][:20] if item['name'] else str(item['code'])))
            ws3.cell(row=i,column=10,value=item['qty'])
        bar3=BarChart(); bar3.type='bar'; bar3.style=10
        bar3.title=f'Top Items by Qty — {sale_date}'
        bar3.y_axis.title='Item'; bar3.x_axis.title='Qty Sold'
        bar3.add_data(Reference(ws3,min_col=10,min_row=2,max_row=2+len(top15q)),titles_from_data=True)
        bar3.set_categories(Reference(ws3,min_col=9,min_row=3,max_row=2+len(top15q)))
        bar3.width=22; bar3.height=14; ws3.add_chart(bar3,'I4')

    # ── Sheet 4: Revenue Ranking ──────────────────────────
    ws4 = wb.create_sheet('Revenue Ranking')
    ws4.freeze_panes = 'A3'
    title_row(ws4, f'{title} — Items Ranked by Revenue (High to Low)', 7, bg='1A3C8F')
    hdr_row(ws4, ['Rank','Item Code','Item Name','Department','Supplier','Total Value (Rs)','Qty Sold'], bg='1A3C8F')
    for ri, item in enumerate(by_value, 3):
        bg = 'F0F4FF' if ri%2==0 else 'FFFFFF'
        vals = [ri-2, item['code'], item['name'], item['dept'], item['supplier'], item['value'], item['qty']]
        for ci,v in enumerate(vals,1):
            c=ws4.cell(row=ri,column=ci,value=v)
            st(c,bold=(ci==1),fg='0D2456' if ci in (1,6) else '1A2340',bg=bg,
               align='center' if ci in (1,2) else 'right' if ci in (6,7) else 'left',
               fmt='#,##0.00' if ci==6 else '#,##0' if ci==7 else '0' if ci==2 else None,sz=9)
    for col,w in zip('ABCDEFG',[7,11,36,24,20,18,11]): ws4.column_dimensions[col].width=w
    # Chart
    top15v = by_value[:min(15,len(by_value))]
    if top15v:
        ws4.cell(row=2,column=9,value='Item'); ws4.cell(row=2,column=10,value='Revenue')
        for i,item in enumerate(top15v,3):
            ws4.cell(row=i,column=9,value=(item['name'][:20] if item['name'] else str(item['code'])))
            ws4.cell(row=i,column=10,value=item['value'])
        bar4=BarChart(); bar4.type='col'; bar4.style=10
        bar4.title=f'Top Items by Revenue — {sale_date}'
        bar4.y_axis.title='Revenue (Rs)'; bar4.x_axis.title='Item'
        bar4.add_data(Reference(ws4,min_col=10,min_row=2,max_row=2+len(top15v)),titles_from_data=True)
        bar4.set_categories(Reference(ws4,min_col=9,min_row=3,max_row=2+len(top15v)))
        bar4.width=22; bar4.height=14; ws4.add_chart(bar4,'I4')

    # ── Sheet 5: Charts ───────────────────────────────────
    ws5 = wb.create_sheet('Charts')
    title_row(ws5, title, 12)
    sup_names=[s['name'] for s in summary.get('suppliers',[])]
    sup_vals=[s['value'] for s in summary.get('suppliers',[])]
    dept_names=[d['name'] for d in summary.get('departments',[])]
    dept_vals=[d['value'] for d in summary.get('departments',[])]
    dept_qtys=[d['qty'] for d in summary.get('departments',[])]
    ws5['A3']='Supplier'; ws5['B3']='Value'
    for i,(n,v) in enumerate(zip(sup_names,sup_vals),4):
        ws5.cell(row=i,column=1,value=n); ws5.cell(row=i,column=2,value=v)
    ws5['D3']='Department'; ws5['E3']='Value'; ws5['F3']='Qty'
    for i,(n,v,q) in enumerate(zip(dept_names,dept_vals,dept_qtys),4):
        ws5.cell(row=i,column=4,value=n); ws5.cell(row=i,column=5,value=v); ws5.cell(row=i,column=6,value=q)
    if sup_names:
        pie5=PieChart(); pie5.title='Supplier-wise Sales'; pie5.style=10
        pie5.add_data(Reference(ws5,min_col=2,min_row=3,max_row=3+len(sup_names)),titles_from_data=True)
        pie5.set_categories(Reference(ws5,min_col=1,min_row=4,max_row=3+len(sup_names)))
        pie5.width=16; pie5.height=13; ws5.add_chart(pie5,'A5')
    if dept_names:
        bar5a=BarChart(); bar5a.type='col'; bar5a.style=10; bar5a.title='Department Revenue'
        bar5a.add_data(Reference(ws5,min_col=5,min_row=3,max_row=3+len(dept_names)),titles_from_data=True)
        bar5a.set_categories(Reference(ws5,min_col=4,min_row=4,max_row=3+len(dept_names)))
        bar5a.width=20; bar5a.height=13; ws5.add_chart(bar5a,'I5')
        bar5b=BarChart(); bar5b.type='col'; bar5b.style=10; bar5b.title='Department Items Sold'
        bar5b.add_data(Reference(ws5,min_col=6,min_row=3,max_row=3+len(dept_names)),titles_from_data=True)
        bar5b.set_categories(Reference(ws5,min_col=4,min_row=4,max_row=3+len(dept_names)))
        bar5b.width=20; bar5b.height=13; ws5.add_chart(bar5b,'I23')

    # ── Sheet 6: Statistics ───────────────────────────────
    ws6 = wb.create_sheet('Statistics')
    title_row(ws6, f'Sales Statistics — {sale_date}', 2)
    hdr_row(ws6, ['Metric','Value'])
    total_cost   = sum(i['cost']  for i in all_items)
    total_profit = sum(i['value'] - i['cost'] for i in all_items)
    profit_pct   = round((total_profit / summary.get('total_value', 1)) * 100, 2) if summary.get('total_value') else 0.0
    stats = [('Total Revenue',      summary.get('total_value', 0), '#,##0.00'),
             ('Total Cost Value',   round(total_cost, 2),          '#,##0.00'),
             ('Total Profit',       round(total_profit, 2),        '#,##0.00'),
             ('Profit %',           profit_pct,                    '0.00"%"'),
             ('Total Items Sold',   summary.get('total_qty', 0),   '#,##0'),
             ('Walk-in Customers',  walkins,                       '#,##0'),
             ('Unique Items',       len(item_map),                 '#,##0'),
             ('Active Suppliers',   len([s for s in summary.get('suppliers',[]) if s['value']>0]), '#,##0'),
             ('Active Departments', len([d for d in summary.get('departments',[]) if d['value']>0]), '#,##0')]
    if all_items:
        stats.append(('Avg Revenue per Item', sum(i['value'] for i in all_items)/len(all_items), '#,##0.00'))
        if by_qty:   stats.append(('Highest Qty Item',     by_qty[0]['name'],   None))
        if by_value: stats.append(('Highest Revenue Item', by_value[0]['name'], None))
    for i,(k,v,fm) in enumerate(stats,3):
        bg='F0F4FF' if i%2==0 else 'FFFFFF'
        # Colour profit rows green/red
        is_profit_row = k in ('Total Profit', 'Profit %')
        val_fg = ('1A7A3A' if (isinstance(v,(int,float)) and v>=0) else 'C0392B') if is_profit_row else '1A3C8F'
        c1=ws6.cell(row=i,column=1,value=k); st(c1,bold=True,fg='1A2340',bg=bg,sz=10)
        c2=ws6.cell(row=i,column=2,value=v); st(c2,fg=val_fg,bg=bg,align='right',fmt=fm,sz=10,bold=True)
    ws6.column_dimensions['A'].width=28; ws6.column_dimensions['B'].width=24

    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True,
                     download_name=f'{sale_date}_SALE_REPORT.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

if __name__=='__main__':
    app.run(debug=True, port=8080)
